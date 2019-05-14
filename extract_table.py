import cv2
import numpy as np
from ctpn.ctpn_blstm_test import text_predict   # Text detection
from densent_ocr.densenet_ocr_test import predict   # Text recognition


def table_lines(src):
    src_height, src_width = src.shape[:2]
    gray = cv2.cvtColor(src, cv2.COLOR_BGR2GRAY)

    erode_size = int(src_height / 300)
    if erode_size % 2 == 0:
        erode_size += 1
    element = cv2.getStructuringElement(cv2.MORPH_RECT, (erode_size, erode_size))
    erod = cv2.erode(gray, element)
    blur_size = int(src_height / 200)
    if blur_size % 2 == 0:
        blur_size += 1
    blur = cv2.GaussianBlur(erod, (blur_size, blur_size), 0, 0)

    thresh = cv2.adaptiveThreshold(~gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY, 15, -2)

    horizontal = thresh
    vertical = thresh

    scale = 20

    print(1111111111111, horizontal.shape)
    horizontalsize = int(horizontal.shape[0] / scale)
    print(222222222, horizontalsize)

    horizontalStructure = cv2.getStructuringElement(cv2.MORPH_RECT, (horizontalsize, 1))
    horizontal = cv2.erode(horizontal, horizontalStructure, (-1, -1))
    horizontal = cv2.dilate(horizontal, horizontalStructure, (-1, -1))
    horizontal = cv2.blur(horizontal, (3, 3))
    # Image.fromarray(horizontal).show()
    # horizontal = cv2.dilate(horizontal, horizontalStructure, (20, 20))

    verticalsize = int(vertical.shape[1] / scale)

    verticalStructure = cv2.getStructuringElement(cv2.MORPH_RECT, (1, verticalsize))
    vertical = cv2.erode(vertical, verticalStructure, (-1, -1))
    vertical = cv2.dilate(vertical, verticalStructure, (-1, -1))
    # Image.fromarray(vertical).show()
    vertical = cv2.blur(vertical, (3, 3))

    joints = cv2.bitwise_and(horizontal, vertical)

    mask = horizontal + vertical

    joints = cv2.bitwise_and(horizontal, vertical)
    # print(horizontal)
    # Image.fromarray(joints).show()
    if not joints.any():
        return 'not table'

    _, contours, hierarchy = cv2.findContours(mask, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
    # cv2.drawContours(src, contours, -1, (0, 255, 0), 1)
    # Image.fromarray(src).show()
    # print(22222222, joints.shape)
    contours_poly = [''] * len(contours)
    boundRect = [''] * len(contours)
    rois = []
    for i in range(len(contours)):
        area = cv2.contourArea(contours[i])
        if area < 100:
            continue
        contours_poly[i] = cv2.approxPolyDP(np.array(contours[i]), 3, True)
        boundRect[i] = cv2.boundingRect(np.array(contours_poly[i]))
        # boundRect[i] = cv2.boundingRect(np.array(contours[i]))

        print(joints.shape, boundRect[i])
        roi = joints[boundRect[i][1]:boundRect[i][1] + boundRect[i][3],
              boundRect[i][0]:boundRect[i][0] + boundRect[i][2]]

        _, joints_contours, _ = cv2.findContours(roi, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
        print(2222222222, len(joints_contours))
        # print(joints_contours)
        if len(joints_contours) < 1:
            continue

        ytt = src[boundRect[i][1]:boundRect[i][1] + boundRect[i][3], boundRect[i][0]:boundRect[i][0] + boundRect[i][2]]

        # rois.append(src(boundRect[i]).clone())
        rois.append([ytt, list(boundRect[i])])

    _, new_con, _ = cv2.findContours(joints, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
    num_x = []
    num_y = []
    for i in new_con:
        num_x.append(cv2.minEnclosingCircle(i)[0][0])
        num_y.append(cv2.minEnclosingCircle(i)[0][1])
    num_x = sorted(num_x)
    num_y = sorted(num_y)
    xs = set()
    for index in range(len(num_x) - 1):
        if abs(num_x[index] - num_x[index + 1]) < 50:
            num_x[index + 1] = num_x[index]
            xs.add(num_x[index])
    ys = set()
    for index in range(len(num_y) - 1):
        if abs(num_y[index] - num_y[index + 1]) < 50:
            num_y[index + 1] = num_y[index]
            ys.add(num_y[index])
    return len(xs) - 1, len(ys) - 1, list(xs), list(ys), rois


def extract_table(image_path):
    from PIL import Image
    src = cv2.imread(image_path)
    if not src.data:
        print('not picture')
    src_height, src_width = src.shape[:2]
    gray = cv2.cvtColor(src, cv2.COLOR_BGR2GRAY)

    erode_size = int(src_height / 300)
    if erode_size % 2 == 0:
        erode_size += 1
    element = cv2.getStructuringElement(cv2.MORPH_RECT, (erode_size, erode_size))
    erod = cv2.erode(gray, element)
    blur_size = int(src_height / 200)
    if blur_size % 2 == 0:
        blur_size += 1
    blur = cv2.GaussianBlur(erod, (blur_size, blur_size), 0, 0)

    thresh = cv2.adaptiveThreshold(~gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY, 15, -2)

    horizontal = thresh
    vertical = thresh

    scale = 20

    print(1111111111111, horizontal.shape)
    horizontalsize = int(horizontal.shape[0] / scale)
    print(222222222, horizontalsize)

    horizontalStructure = cv2.getStructuringElement(cv2.MORPH_RECT, (horizontalsize, 1))
    horizontal = cv2.erode(horizontal, horizontalStructure, (-1, -1))
    horizontal = cv2.dilate(horizontal, horizontalStructure, (-1, -1))
    horizontal = cv2.blur(horizontal, (3, 3))
    # Image.fromarray(horizontal).show()
    # horizontal = cv2.dilate(horizontal, horizontalStructure, (20, 20))

    verticalsize = int(vertical.shape[1] / scale)

    verticalStructure = cv2.getStructuringElement(cv2.MORPH_RECT, (1, verticalsize))
    vertical = cv2.erode(vertical, verticalStructure, (-1, -1))
    vertical = cv2.dilate(vertical, verticalStructure, (-1, -1))
    # Image.fromarray(vertical).show()
    vertical = cv2.blur(vertical, (3, 3))

    mask = horizontal + vertical

    joints = cv2.bitwise_and(horizontal, vertical)
    # print(horizontal)
    # Image.fromarray(joints).show()
    if not joints.any():
        return 'not table'

    _, contours, hierarchy = cv2.findContours(mask, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
    # cv2.drawContours(src, contours, -1, (0, 255, 0), 1)
    # Image.fromarray(src).show()
    # print(22222222, joints.shape)
    contours_poly = [''] * len(contours)
    boundRect = [''] * len(contours)
    rois = []
    for i in range(len(contours)):
        area = cv2.contourArea(contours[i])
        if area < 100:
            continue
        contours_poly[i] = cv2.approxPolyDP(np.array(contours[i]), 3, True)
        boundRect[i] = cv2.boundingRect(np.array(contours_poly[i]))
        # boundRect[i] = cv2.boundingRect(np.array(contours[i]))

        print(joints.shape, boundRect[i])
        roi = joints[boundRect[i][1]:boundRect[i][1] + boundRect[i][3],
              boundRect[i][0]:boundRect[i][0] + boundRect[i][2]]

        _, joints_contours, _ = cv2.findContours(roi, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
        print(2222222222, len(joints_contours))
        # print(joints_contours)
        if len(joints_contours) < 1:
            continue

        ytt = src[boundRect[i][1]:boundRect[i][1] + boundRect[i][3], boundRect[i][0]:boundRect[i][0] + boundRect[i][2]]
        # Image.fromarray(ytt).save('data/{}.jpg'.format(i))

        # rois.append(src(boundRect[i]).clone())
        rois.append([ytt, list(boundRect[i])])
        # cv2.rectangle(src, (boundRect[i][0], boundRect[i][1]), (boundRect[i][0] + boundRect[i][2], boundRect[i][1] + boundRect[i][3]), (0, 255, 0), 3)
    new_img = sorted(rois, key=lambda i: i[1][3])[-1][0]

    cols, rows, col_point, row_point, tables = table_lines(new_img)
    col_point = sorted(col_point)
    row_point = sorted(row_point)
    print(col_point)
    print(row_point)
    tables = sorted(tables, key=lambda i: i[1][3])[:-1]
    tables = sorted(tables, key=lambda i: i[1][0] + i[1][1])
    # print(tables[14][1])
    # Image.fromarray(tables[14][0]).show()
    for i in tables:
        d = {'col_begin': 0, 'col_end': 0, 'row_begin': 0, 'row_end': 0}
        for index, value in enumerate(col_point):
            if index == 0:
                d_range = 50
            else:
                d_range = (col_point[index] - col_point[index - 1]) / 2
            if i[1][0] > col_point[index] - d_range:
                # print(33333333333, i[1], index)
                d['col_begin'] = index
        for index, value in enumerate(col_point):
            if index == len(col_point) - 1:
                d_range = 50
            else:
                d_range = (col_point[index + 1] - col_point[index]) / 2
            if i[1][0] + i[1][2] < col_point[index] + d_range:
                d['col_end'] = index
                break
        for index, value in enumerate(row_point):
            if index == 0:
                d_range = 50
            else:
                d_range = (row_point[index] - row_point[index - 1]) / 2
            if i[1][1] > row_point[index] - d_range:
                d['row_begin'] = index
        for index, value in enumerate(row_point):
            if index == len(row_point) - 1:
                d_range = 50
            else:
                d_range = (row_point[index + 1] - row_point[index]) / 2
            if i[1][1] + i[1][3] < row_point[index] + d_range:
                d['row_end'] = index
                break
        i.append(d)
    # for i in tables:
    #     print(i[2])
    # Image.fromarray(new_img[111: 111+209, 344:344+319]).show()
    # print(44444444, cols, rows, tables)
    generate_table(cols, rows, tables)
    Image.fromarray(src).save('space_d.jpg')


import docx


def generate_table(cols, rows, tables):
    from PIL import Image
    document = docx.Document()
    table = document.add_table(rows, cols)
    for i in tables:
        d = i[2]
        # if d['col_end'] - d['col_begin'] == 1 and d['row_end'] - d['row_begin'] == 1:
        #     continue
        if d['col_end'] - d['col_begin'] != 1:
            for col in range(d['col_begin'], d['col_end']):
                table.cell(d['row_begin'], d['col_begin']).merge(table.cell(d['row_begin'], col))
        if d['row_end'] - d['row_begin'] != 1:
            for row in range(d['row_begin'], d['row_end']):
                table.cell(d['row_begin'], d['col_begin']).merge(table.cell(row, d['col_begin']))
        texts = ''
        images = text_predict(i[0])
        if not images:
            texts += predict(i[0])[0]
        for image in sorted(images, key=lambda i:i[0][1]):
            if image[1].any():
                texts += predict(image[1])[0]
        print(texts)
        table.cell(d['row_begin'], d['col_begin']).text = texts
    document.save('space.docx')


extract_table('04.jpg')
